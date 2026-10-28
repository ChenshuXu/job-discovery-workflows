#!/usr/bin/env python3
"""Official PERM local history: explicit refresh/import, read-only offline query/preview."""
import argparse
import contextlib
import datetime as dt
import html.parser
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import statistics
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
import uuid
import zipfile

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / '.local/perm-history'
INDEX = 'https://www.dol.gov/agencies/eta/foreign-labor/performance'
VERSION = 1
STATUSES = ('Certified', 'Certified-Expired', 'Denied', 'Withdrawn')
FIELDS = {
 'case_number': ('CASE_NUMBER',), 'status': ('CASE_STATUS',),
 'received': ('RECEIVED_DATE',), 'decision': ('DECISION_DATE',),
 'employer': ('EMP_BUSINESS_NAME', 'EMPLOYER_NAME'), 'trade': ('EMP_TRADE_NAME', 'EMPLOYER_TRADE_NAME'),
 'employer_city': ('EMP_CITY', 'EMPLOYER_CITY'), 'employer_state': ('EMP_STATE', 'EMPLOYER_STATE_PROVINCE', 'EMPLOYER_STATE'),
 'job_title': ('JOB_TITLE',), 'soc_code': ('PWD_SOC_CODE', 'PW_SOC_CODE'), 'soc_title': ('PWD_SOC_TITLE', 'PW_SOC_TITLE'),
 'worksite_city': ('PRIMARY_WORKSITE_CITY', 'WORKSITE_CITY'), 'worksite_state': ('PRIMARY_WORKSITE_STATE', 'WORKSITE_STATE'),
 'multiple_worksites': ('IS_MULTIPLE_LOCATIONS', 'WORKSITE_MULTIPLE_LOCATIONS', 'WORKSITE_MULTIPLE_LOCATIONS_Y_N', 'MULTIPLE_WORKSITE_LOCATIONS'),
 'law_firm': ('ATTY_AG_LAW_FIRM_NAME', 'AGENT_ATTORNEY_FIRM_NAME'),
 'representative_last': ('ATTY_AG_LAST_NAME', 'AGENT_ATTORNEY_LAST_NAME'),
 'representative_first': ('ATTY_AG_FIRST_NAME', 'AGENT_ATTORNEY_FIRST_NAME'),
 'representative_middle': ('ATTY_AG_MIDDLE_NAME', 'AGENT_ATTORNEY_MIDDLE_NAME')}
REQUIRED = ('case_number', 'status', 'received', 'decision', 'employer', 'job_title', 'soc_code', 'worksite_city', 'worksite_state', 'law_firm')
BOUNDARY = '官方披露不能完整提供全部申请量及当前待审量；历史认证不证明当前岗位支持 PERM，也不是个人获批概率。'


def without_checksums(value):
    if isinstance(value,dict):
        return {k:without_checksums(v) for k,v in value.items() if 'sha256' not in k.lower() and 'sha-256' not in k.lower()}
    if isinstance(value,list):return [without_checksums(v) for v in value]
    return value


def guarded(path, *, inside=True):
    path = Path(os.path.abspath(path))
    for part in [path, *path.parents]:
        if part.is_symlink(): raise ValueError(f'Symlinks are forbidden: {part}')
    if inside and not path.is_relative_to(DATA): raise ValueError('Data path is outside checkout .local/perm-history')
    return path


def write_json(path, value):
    guarded(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    try:
        with temp.open('x') as f:
            json.dump(without_checksums(value), f, ensure_ascii=False, indent=2); f.write('\n'); f.flush(); os.fsync(f.fileno())
        os.replace(temp, path)
    finally:
        if temp.exists(): temp.unlink()


def read_json(path, inside=True):
    return without_checksums(json.loads(guarded(path, inside=inside).read_text()))


@contextlib.contextmanager
def writer():
    guarded(DATA); DATA.mkdir(parents=True, exist_ok=True)
    lock = guarded(DATA / 'write.lock')
    try: fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError: raise ValueError('PERM writer is locked; inspect existing write.lock before retrying')
    os.write(fd, json.dumps({'pid':os.getpid(),'started':now()}).encode()); os.close(fd)
    try: yield
    finally: lock.unlink()


def now(): return dt.datetime.now(dt.timezone.utc).isoformat()

def norm(value):
    value = unicodedata.normalize('NFKC', str(value or '')).casefold()
    return ' '.join(''.join(c if c.isalnum() else ' ' for c in value).split())


def stem(value):
    words = norm(value).split()
    while words and words[-1] in {'inc','incorporated','llc','llp','lp','ltd','limited','corporation','corp','company','co','pllc'}: words.pop()
    return ' '.join(words)


def official(url):
    p = urllib.parse.urlsplit(url)
    if p.scheme != 'https' or not p.hostname or not (p.hostname == 'dol.gov' or p.hostname.endswith('.dol.gov')) or p.username or p.password or p.port not in (None,443):
        raise ValueError('Only official HTTPS dol.gov URLs are permitted')
    return url


class OfficialRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        official(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch(url, dest=None):
    official(url)
    # DOL rejects the custom agent string; use its verified working curl client header.
    request = urllib.request.Request(url, headers={'User-Agent':'curl/8.7.1'})
    with urllib.request.build_opener(OfficialRedirect()).open(request, timeout=120) as r:
        official(r.url)
        if r.status != 200: raise ValueError(f'Unexpected HTTP status {r.status}')
        expected = int(r.headers['Content-Length']) if r.headers.get('Content-Length') else None
        if dest:
            dest = guarded(dest); total = 0
            with dest.open('xb') as f:
                while block := r.read(1024 * 1024): f.write(block); total += len(block)
            if expected is not None and total != expected: raise ValueError('Truncated official download')
            return {'url':url,'effective_url':r.url,'path':str(dest),'bytes':total,'fetched_at':now()}
        body = r.read()
        if expected is not None and len(body) != expected: raise ValueError('Truncated official index')
        return body.decode('utf-8')


class Links(html.parser.HTMLParser):
    def __init__(self): super().__init__(); self.links=[]
    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            for k,v in attrs:
                if k == 'href' and v: self.links.append(urllib.parse.urljoin(INDEX,v))


def discover():
    page = fetch(INDEX); links = Links(); links.feed(page)
    result=[]
    for fy in (2025,2026):
        xlsx=[]
        for url in links.links:
            match=re.search(r'PERM_Disclosure_Data_FY'+str(fy)+r'_Q([1-4])\.xlsx(?:\?|$)',url,re.I)
            if match: xlsx.append((int(match.group(1)),official(url)))
        if not xlsx: raise ValueError(f'Official index has no FY{fy} PERM disclosure')
        q=max(x[0] for x in xlsx)
        if fy == 2025 and q != 4: raise ValueError('FY2025 full-year Q4 is required')
        candidates=sorted({u for quarter,u in xlsx if quarter==q})
        if len(candidates)!=1: raise ValueError('Ambiguous official release URLs')
        layouts=sorted({official(u) for u in links.links if re.search(r'PERM_Record_Layout_FY'+str(fy)+r'_Q'+str(q)+r'\.pdf(?:\?|$)',u,re.I)})
        if len(layouts)!=1: raise ValueError(f'No unique official FY{fy} Q{q} layout')
        end=(dt.date(fy-1,12,31),dt.date(fy,3,31),dt.date(fy,6,30),dt.date(fy,9,30))[q-1]
        result.append({'fiscal_year':fy,'quarter':q,'coverage_start':f'{fy-1}-10-01','coverage_end':str(end),'xlsx_url':candidates[0],'layout_url':layouts[0]})
    return page,result


def validate_source_manifest(manifest):
    if manifest.get('schema_version') != 1: raise ValueError('Unsupported source manifest')
    sources=manifest.get('sources',[])
    if len(sources)!=2 or sorted(s.get('fiscal_year') for s in sources)!=[2025,2026]: raise ValueError('Exactly one FY2025 and one FY2026 source required; cumulative quarters cannot be combined')
    for s in sources:
        fy=s['fiscal_year'];q=s['quarter']
        if q not in (1,2,3,4) or (fy==2025 and q!=4): raise ValueError('Invalid quarter/full-year selection')
        end=[f'{fy-1}-12-31',f'{fy}-03-31',f'{fy}-06-30',f'{fy}-09-30'][q-1]
        if s['coverage_start']!=f'{fy-1}-10-01' or s['coverage_end']!=end: raise ValueError('Coverage dates do not match release')
        for key in ('xlsx','layout'):
            info=s[key];official(info['url']);f=guarded(info['path'])
            if not f.is_file(): raise ValueError(f'Source file is missing: {f}')
            if info.get('bytes') is not None and f.stat().st_size != info['bytes']: raise ValueError('Source byte length mismatch')
            with f.open('rb') as stream: signature=stream.read(8)
            if key=='xlsx' and signature[:4]!=b'PK\x03\x04': raise ValueError('Not an XLSX ZIP file')
            if key=='layout' and not signature.startswith(b'%PDF-'): raise ValueError('Not a PDF layout')
        # ZIP CRC validation detects truncated/corrupt files before parsing any data.
        with zipfile.ZipFile(s['xlsx']['path']) as z:
            if '[Content_Types].xml' not in z.namelist(): raise ValueError('Invalid XLSX structure')
            if z.testzip(): raise ValueError('XLSX CRC validation failed')
    return sorted(sources,key=lambda x:(x['fiscal_year'],x['quarter']))


def date_value(value):
    if value in (None,''): return None,'missing'
    if isinstance(value,dt.datetime): return value.date().isoformat(),None
    if isinstance(value,dt.date): return value.isoformat(),None
    value=str(value).strip()
    for fmt in ('%Y-%m-%d','%m/%d/%Y','%m/%d/%y','%Y-%m-%d %H:%M:%S'):
        try:return dt.datetime.strptime(value,fmt).date().isoformat(),None
        except ValueError:pass
    return None,'invalid'


def import_sources(manifest):
    started=time.monotonic(); sources=validate_source_manifest(manifest)
    from openpyxl import load_workbook
    revision=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'-'+uuid.uuid4().hex[:8]
    stage=guarded(DATA/('.staging-'+revision));stage.mkdir()
    audit={'schema_version':1,'started_at':now(),'rows_read':0,'deduplicated_rows':0,'cross_release_revisions':0,'sources':[],'date_anomalies':{'missing_received':0,'invalid_received':0,'missing_decision':0,'invalid_decision':0,'negative_duration':0}}
    old=guarded(DATA/'current.json').read_bytes() if (DATA/'current.json').exists() else None
    db=None
    try:
        (stage/'raw').mkdir();(stage/'layouts').mkdir()
        db=sqlite3.connect(stage/'cases.sqlite')
        db.executescript('CREATE TABLE cases(case_number TEXT PRIMARY KEY, employer_norm TEXT NOT NULL, trade_norm TEXT NOT NULL, decision TEXT, release_order INTEGER NOT NULL, payload TEXT NOT NULL); CREATE INDEX employer_lookup ON cases(employer_norm); CREATE INDEX trade_lookup ON cases(trade_norm); CREATE TABLE case_versions(case_number TEXT NOT NULL, decision TEXT, release_order INTEGER NOT NULL, payload TEXT NOT NULL); CREATE INDEX case_revision_lookup ON case_versions(case_number,decision,release_order);')
        for source in sources:
            s=json.loads(json.dumps(source));fy=s['fiscal_year'];q=s['quarter'];order=fy*10+q
            for key,folder in [('xlsx','raw'),('layout','layouts')]:
                src=guarded(s[key]['path']);name=f'FY{fy}_Q{q}'+('.xlsx' if key=='xlsx' else '.pdf');dest=stage/folder/name
                shutil.copyfile(src,dest);s[key]['path']=f'{folder}/{name}';s[key]['bytes']=dest.stat().st_size
            book=load_workbook(stage/s['xlsx']['path'],read_only=True,data_only=True)
            source_audit={'fiscal_year':fy,'quarter':q,'sheets':[],'rows':0,'status_counts':{x:0 for x in STATUSES}}
            try:
                for sheet in book:
                    rows=sheet.iter_rows(values_only=True);mapping=None;header=None;header_row=None
                    for line,values in enumerate(rows,1):
                        candidates=[str(v or '').strip().upper() for v in values]
                        if 'CASE_NUMBER' in candidates:
                            if len([x for x in candidates if x])!=len(set(x for x in candidates if x)): raise ValueError('Duplicate worksheet headers')
                            header=candidates;header_row=line;mapping={field:next((candidates.index(alias) for alias in aliases if alias in candidates),None) for field,aliases in FIELDS.items()};break
                        if line>=20:break
                    if mapping is None:continue
                    if any(mapping[x] is None for x in REQUIRED):raise ValueError(f'Missing required columns in {sheet.title}: '+','.join(x for x in REQUIRED if mapping[x] is None))
                    sheet_audit={'sheet':sheet.title,'header_row':header_row,'headers':header,'mapping':{k:header[v] if v is not None else None for k,v in mapping.items()},'rows':0}
                    source_audit['sheets'].append(sheet_audit)
                    for line,values in enumerate(rows,header_row+1):
                        if not any(v not in (None,'') for v in values):continue
                        record={k:(str(values[i]).strip() if i is not None and i<len(values) and values[i] is not None else '') for k,i in mapping.items()}
                        if not record['case_number'] or not record['employer']:raise ValueError('Missing case number or employer')
                        record['status_raw']=record['status']
                        status=re.sub(r'\s*-\s*','-',record['status'].casefold().replace('_','-'))
                        record['status']=next((x for x in STATUSES if x.casefold()==status),None)
                        if record['status'] is None:raise ValueError(f'Unknown case status {status}')
                        for field in ('received','decision'):
                            original=values[mapping[field]];record[field],error=date_value(original)
                            record[field+'_raw']=str(original or '')
                            record[field+'_error']=error
                            if error:audit['date_anomalies'][error+'_'+field]+=1
                        record['processing_days']=None
                        if record['received'] and record['decision']:
                            days=(dt.date.fromisoformat(record['decision'])-dt.date.fromisoformat(record['received'])).days
                            if days>=0:record['processing_days']=days
                            else:audit['date_anomalies']['negative_duration']+=1
                        if record['decision'] and not source['coverage_start']<=record['decision']<=source['coverage_end']:raise ValueError('Decision date outside official source coverage')
                        record.update({'source_fiscal_year':fy,'source_quarter':q,'decision_fiscal_year':(dt.date.fromisoformat(record['decision']).year+(dt.date.fromisoformat(record['decision']).month>=10)) if record['decision'] else fy,'source_file':s['xlsx']['path'],'source_url':s['xlsx']['url'],'sheet':sheet.title,'row':line})
                        payload=json.dumps(record,sort_keys=True,ensure_ascii=False)
                        same_release=db.execute('SELECT payload FROM case_versions WHERE case_number=? AND decision IS ? AND release_order=? LIMIT 1',(record['case_number'],record['decision'],order)).fetchone()
                        ignored={'sheet','row','source_file','source_url','source_fiscal_year','source_quarter'}
                        if same_release and {k:v for k,v in json.loads(same_release[0]).items() if k not in ignored}!={k:v for k,v in record.items() if k not in ignored}:raise ValueError('Conflicting same-release/date case: '+record['case_number'])
                        existing=db.execute('SELECT decision,release_order,payload FROM cases WHERE case_number=?',(record['case_number'],)).fetchone()
                        replace=True
                        if existing:
                            prior=json.loads(existing[2]);same_order=existing[1]==order;same_date=existing[0]==record['decision']
                            ignored={'sheet','row','source_file','source_url','source_fiscal_year','source_quarter'}
                            if same_order and same_date and {k:v for k,v in prior.items() if k not in ignored}!={k:v for k,v in record.items() if k not in ignored}:raise ValueError('Conflicting same-release/date case: '+record['case_number'])
                            audit['deduplicated_rows']+=1
                            if existing[1]!=order:audit['cross_release_revisions']+=1
                            replace=(record['decision'] or '',order)>(existing[0] or '',existing[1])
                        db.execute('INSERT INTO case_versions VALUES (?,?,?,?)',(record['case_number'],record['decision'],order,payload))
                        if replace:db.execute('INSERT OR REPLACE INTO cases VALUES (?,?,?,?,?,?)',(record['case_number'],norm(record['employer']),norm(record['trade']),record['decision'],order,payload))
                        audit['rows_read']+=1;source_audit['rows']+=1;sheet_audit['rows']+=1;source_audit['status_counts'][record['status']]+=1
                        if audit['rows_read']%10000==0:db.commit()
            finally:book.close()
            if source_audit['rows']==0:raise ValueError('No cases in source workbook')
            audit['sources'].append(source_audit)
            source.clear();source.update(s)
        db.commit();audit['unique_cases']=db.execute('SELECT COUNT(*) FROM cases').fetchone()[0]
        if db.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise ValueError('SQLite integrity check failed')
        db.close();db=None
        import resource
        audit.update({'finished_at':now(),'elapsed_seconds':round(time.monotonic()-started,3),'peak_rss_platform_units':resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,'peak_rss_unit':'bytes on macOS; KiB on Linux','status':'COMPLETE'})
        write_json(stage/'import-audit.json',audit)
        output={'schema_version':1,'parser_version':VERSION,'revision':revision,'created_at':now(),'official_index_url':manifest.get('official_index_url',INDEX),'sources':sources,'database':{'path':'cases.sqlite'},'unique_cases':audit['unique_cases'],'boundary':BOUNDARY}
        write_json(stage/'manifest.json',output)
        current=guarded(DATA/'current.json').read_bytes() if (DATA/'current.json').exists() else None
        if current!=old:raise ValueError('Current pointer changed during import')
        versions=guarded(DATA/'versions');versions.mkdir(exist_ok=True);final=versions/revision;os.replace(stage,final)
        for f in final.rglob('*'):
            if f.is_file():f.chmod(0o444)
        pointer={'schema_version':1,'revision':revision,'path':f'versions/{revision}'}
        write_json(DATA/'current.json',pointer)
        return {'status':'COMPLETE','revision':revision,'directory':str(final),'unique_cases':audit['unique_cases'],'rows_read':audit['rows_read'],'elapsed_seconds':audit['elapsed_seconds']}
    except Exception as e:
        if db:db.close()
        audit.update({'status':'FAILED','error':str(e),'finished_at':now()})
        write_json(stage/'failure.json',audit)
        raise


def load_current(pointer=None):
    pointer=pointer if pointer is not None else read_json(DATA/'current.json')
    if pointer.get('schema_version')!=1:raise ValueError('Unsupported current pointer')
    directory=guarded(DATA/pointer['path'])
    if directory.parent!=DATA/'versions' or directory.name!=pointer['revision']:raise ValueError('Invalid version pointer')
    manifest_path=guarded(directory/'manifest.json')
    manifest=read_json(manifest_path)
    if manifest.get('schema_version')!=1 or manifest.get('revision')!=pointer['revision']:raise ValueError('Dataset manifest/version mismatch')
    for source in manifest['sources']:
        for key in ('xlsx','layout'):
            evidence=guarded(directory/source[key]['path'])
            if not evidence.is_relative_to(directory) or not evidence.is_file():raise ValueError('Raw source/layout is missing or outside its version')
            with evidence.open('rb') as f:signature=f.read(8)
            if key=='xlsx' and signature[:4]!=b'PK\x03\x04':raise ValueError('Not an XLSX ZIP file')
            if key=='layout' and not signature.startswith(b'%PDF-'):raise ValueError('Not a PDF layout')
    dbpath=guarded(directory/manifest['database']['path'])
    if not dbpath.is_relative_to(directory) or not dbpath.is_file():raise ValueError('Database is missing or outside its version')
    if read_json(directory/'import-audit.json').get('status')!='COMPLETE':raise ValueError('Dataset import is incomplete')
    uri=dbpath.as_uri()+'?mode=ro&immutable=1'
    return sqlite3.connect(uri,uri=True),manifest,pointer


def quantile(values,p):
    if not values:return None
    position=(len(values)-1)*p;lower=int(position);upper=min(lower+1,len(values)-1)
    return values[lower]+(values[upper]-values[lower])*(position-lower)


def summarize(records):
    counts={s:sum(r['status']==s for r in records) for s in STATUSES};n=counts['Certified']+counts['Certified-Expired'];d=n+counts['Denied']
    processing={}
    for status in STATUSES:
        subset=[r for r in records if r['status']==status];values=sorted(r['processing_days'] for r in subset if r['processing_days'] is not None)
        processing[status]={'n':len(values),'excluded':len(subset)-len(values),'median':statistics.median(values) if values else None,'p25':quantile(values,.25),'p75':quantile(values,.75)}
    return {'disclosed_cases':len(records),'status_counts':counts,'historical_certification_share':{'numerator':n,'denominator':d,'value':n/d if d else None},'processing_days':processing,'recent_dates':{'decision':max((r['decision'] for r in records if r['decision']),default=None),'certification_decision':max((r['decision'] for r in records if r['decision'] and r['status'] in ('Certified','Certified-Expired')),default=None),'received_among_disclosed':max((r['received'] for r in records if r['received']),default=None)}}


def validate_aliases(registry,db,evidence_snapshots=None):
    if registry.get('schema_version')!=1 or not isinstance(registry.get('aliases'),list):raise ValueError('Unsupported alias registry')
    seen=set();ids=set()
    for alias in registry['aliases']:
        name=norm(alias.get('source_company'));ident=alias.get('id')
        if not name or name in seen or not isinstance(ident,str) or not ident or ident in ids:raise ValueError('Duplicate/empty alias source or id')
        seen.add(name);ids.add(ident)
        if alias.get('status') not in ('verified','unresolved') or not alias.get('relationship_type'):raise ValueError('Invalid alias approval/relationship')
        evidence=alias.get('evidence',[]);eids=set()
        if not evidence:raise ValueError('Alias evidence is required')
        for e in evidence:
            eid=e.get('id');url=urllib.parse.urlsplit(e.get('url',''))
            if not eid or eid in eids:raise ValueError('Duplicate/empty alias evidence id')
            eids.add(eid)
            if url.scheme!='https' or not url.hostname or url.username or url.password:raise ValueError('Alias evidence needs HTTPS source URL')
            relative=Path(e.get('path',''))
            if relative.is_absolute() or not relative.parts or '..' in relative.parts:raise ValueError('Alias evidence path must be relative within data root')
            if evidence_snapshots is None:
                f=guarded(DATA/relative)
                if not f.is_file():raise ValueError('Alias evidence file is missing')
                text=f.read_text()
            else:
                text=evidence_snapshots.get(e['path'])
                if not isinstance(text,str):raise ValueError('Frozen alias evidence is missing')
            if not isinstance(e.get('quote'),str) or not e['quote'].strip() or e['quote'] not in text:raise ValueError('Alias evidence quote is absent from snapshot')
            try:
                checked=dt.datetime.fromisoformat(e['checked_at'].replace('Z','+00:00'))
                if checked.tzinfo is None:raise ValueError('timezone required')
            except (KeyError,TypeError,ValueError):raise ValueError('Alias evidence needs checked_at with timezone')
        if not alias.get('source_evidence_ids') or not set(alias['source_evidence_ids'])<=eids:raise ValueError('Source company/JD evidence is required')
        scope=alias.get('scope',{})
        company_history=alias['relationship_type']=='company_history_reference'
        if company_history and scope.get('type')!='company':raise ValueError('Company history requires company identity scope')
        if scope.get('type')=='posting':
            postings=scope.get('postings',[])
            pairs={(p.get('posting_key'),p.get('posting_url')) for p in postings}
            if not postings or len(pairs)!=len(postings) or any(not k or not u or urllib.parse.urlsplit(u).scheme!='https' for k,u in pairs):raise ValueError('Alias posting scope needs unique exact key/URL pairs')
            source_urls={e['url'] for e in evidence if e['id'] in alias['source_evidence_ids']}
            if not {u for _,u in pairs}<=source_urls:raise ValueError('Every scoped posting needs its exact source JD evidence URL')
        elif scope.get('type')=='company':
            if scope.get('official_identity_url') not in {e['url'] for e in evidence}:raise ValueError('Company alias scope requires official identity evidence URL')
        else:raise ValueError('Invalid alias scope')
        targets=alias.get('targets',[]);names=set()
        if alias['status']=='verified' and not targets:raise ValueError('Verified alias needs target entities')
        if alias['status']=='unresolved' and (targets or not alias.get('reason')):raise ValueError('Unresolved alias requires reason and no trusted targets')
        for target in targets:
            name=norm(target.get('legal_name'));locations=target.get('locations',[])
            if not name or name in names:raise ValueError('Duplicate/empty alias target legal name')
            names.add(name)
            if not target.get('evidence_ids') or not set(target['evidence_ids'])<=eids:raise ValueError('Alias target evidence is missing')
            pairs={(norm(p.get('city')),norm(p.get('state'))) for p in locations}
            existing={(norm(r['employer_city']),norm(r['employer_state'])) for row in db.execute('SELECT payload FROM cases WHERE employer_norm=?',(name,)) for r in [json.loads(row[0])]}
            if (not pairs and (not company_history or existing)) or len(pairs)!=len(locations) or any(not c or not s for c,s in pairs):raise ValueError('Alias target needs explicit unique city/state locations when disclosed records exist')
            if existing and not pairs<=existing:raise ValueError('Alias target location absent from current PERM legal-name records')
    return registry


def load_aliases(db):
    path=guarded(DATA/'employer-aliases.json')
    if not path.exists():return None
    return {'registry':validate_aliases(read_json(path),db)}


def import_aliases(path,merge=False):
    target=guarded(DATA/'employer-aliases.json');old=target.read_bytes() if target.exists() else None
    db,_,_=load_current()
    try:
        registry=validate_aliases(read_json(path),db)
        if merge:
            previous=validate_aliases(json.loads(old),db) if old else {'schema_version':1,'aliases':[]}
            aliases={norm(a['source_company']):a for a in previous['aliases']}
            aliases.update({norm(a['source_company']):a for a in registry['aliases']})
            registry=validate_aliases({'schema_version':1,'aliases':list(aliases.values())},db)
    finally:db.close()
    if (target.read_bytes() if target.exists() else None)!=old:raise ValueError('Alias registry changed during import')
    write_json(target,registry)
    return {'status':'COMPLETE','path':str(target),'alias_count':len(registry['aliases'])}


def add_statistics(result,records):
    result.update(summarize(records));result['total_disclosed_cases']=result.pop('disclosed_cases')
    result['years']={str(fy):summarize([r for r in records if r['decision_fiscal_year']==fy]) for fy in (2025,2026)}
    result['breakdowns']={}
    for label,fields in [('job_titles',('job_title',)),('soc',('soc_code','soc_title')),('primary_worksites',('worksite_city','worksite_state','multiple_worksites')),('law_firms',('law_firm',)),('representatives',('law_firm','representative_last','representative_first','representative_middle'))]:
        groups={}
        for r in records:
            key=tuple(r[k] or '(unknown)' for k in fields);groups.setdefault(key,[]).append(r)
        result['breakdowns'][label]=[{'values':dict(zip(fields,key)),**summarize(items)} for key,items in sorted(groups.items(),key=lambda x:(-len(x[1]),x[0]))]
    result['cases']=records
    result['worksite_boundary']='Primary worksite only; employer address is not the worksite. Multi-location flag retained when disclosed.'
    return result


def apply_alias(db,result,alias,aliases,contexts):
    result.update({'match_status':alias['status'],'match_method':'evidence_backed_alias','match_reason':alias.get('reason'),'suggestions':[],'matched_employers':[],'alias':alias,'posting_context':contexts})
    if alias['scope']['type']=='posting':
        allowed={(p['posting_key'],p['posting_url']) for p in alias['scope']['postings']}
        if not contexts or any(not p.get('posting_key') or not p.get('posting_url') for p in contexts):
            result.update(match_status='missing_context',match_reason='This alias requires the verified posting key and URL');return result
        if any((p['posting_key'],p['posting_url']) not in allowed for p in contexts):
            result.update(match_status='context_mismatch',match_reason='Posting is outside the verified alias scope');return result
    if alias['status']!='verified':return result
    result['match_status']='verified_alias';result['alias_verification']='verified';entities=[]
    for target in alias['targets']:
        locations={(norm(p['city']),norm(p['state'])) for p in target['locations']}
        records=[json.loads(row[0]) for row in db.execute('SELECT payload FROM cases WHERE employer_norm=? ORDER BY case_number',(norm(target['legal_name']),))]
        records=[r for r in records if (norm(r['employer_city']),norm(r['employer_state'])) in locations]
        entity={'legal_name':target['legal_name'],'locations':target['locations'],'evidence_ids':target['evidence_ids']}
        entity['match_status']='verified_alias' if records else 'not_found'
        if records:add_statistics(entity,records)
        else:entity['match_reason']='Verified legal identity has no disclosed records in the selected local releases; complete application count remains unavailable'
        entities.append(entity)
        result['matched_employers'] += [{'name':target['legal_name'],**p} for p in target['locations']]
    result['entities']=entities
    if len(entities)==1:
        if entities[0].get('cases'):add_statistics(result,entities[0]['cases'])
        else:result.update(match_status='not_found',match_reason=entities[0]['match_reason'])
    else:result['match_reason']='Multiple verified legal entities are displayed separately; no group total is inferred'
    return result


def query(db,manifest,pointer,company,aliases=None,contexts=None,suggestion_cache=None):
    name=norm(company)
    if not name:raise ValueError('Company name is empty')
    legal_records=[json.loads(row[0]) for row in db.execute('SELECT payload FROM cases WHERE employer_norm=? ORDER BY case_number',(name,))]
    records=legal_records or [json.loads(row[0]) for row in db.execute('SELECT payload FROM cases WHERE trade_norm=? ORDER BY case_number',(name,))]
    identity_groups={}
    for r in records:
        raw=(r['employer'],r['employer_city'],r['employer_state'])
        identity_groups.setdefault(tuple(norm(x) for x in raw),set()).add(raw)
    identities=[sorted(identity_groups[key])[0] for key in sorted(identity_groups)]
    result={'schema_version':1,'metrics_version':1,'matching_policy':'normalized exact legal name and location; trade-only names are unconfirmed','company':company,'match_status':'exact' if len(identities)==1 and legal_records else 'ambiguous' if identities else 'not_found','revision':pointer['revision'],'coverage':[{'fiscal_year':s['fiscal_year'],'quarter':s['quarter'],'start':s['coverage_start'],'end':s['coverage_end']} for s in manifest['sources']],'matched_employers':[{'name':e,'city':c,'state':s,'raw_variants':[{'name':a,'city':b,'state':c2} for a,b,c2 in sorted(identity_groups[(norm(e),norm(c),norm(s))])]} for e,c,s in identities],'suggestions':[],'complete_application_count':None,'pending_count':None,'boundary':BOUNDARY}
    if aliases:
        alias=next((a for a in aliases['registry']['aliases'] if norm(a['source_company'])==name),None)
        if alias:return apply_alias(db,result,alias,aliases,contexts or [])
    result['match_method']='exact_legal_name' if legal_records else 'unconfirmed_trade_name' if identities else None
    if identities and not legal_records:
        result['match_reason']='Exact trade name alone does not bind the scored company to this legal entity; independent alias evidence is required'
    elif len(identities)>1:
        result['match_reason']='Exact legal name has conflicting employer locations; identity requires disambiguation'
    if not identities:
        pairs=db.execute('SELECT DISTINCT json_extract(payload,"$.employer"),json_extract(payload,"$.trade") FROM cases')
        if suggestion_cache is not None:
            if 'index' not in suggestion_cache:
                index={}
                for employer,trade in pairs:
                    for suffix in {stem(employer),stem(trade)}:
                        if suffix:index.setdefault(suffix,[]).append((employer,trade))
                suggestion_cache['index']=index
            pairs=suggestion_cache['index'].get(stem(company),[])
        for employer,trade in pairs:
            if stem(company) and stem(company) in {stem(employer),stem(trade)}:result['suggestions'].append({'employer':employer,'trade':trade,'reason':'legal-suffix variant only; requires independent identity evidence'})
    if result['match_status']!='exact':return result
    return add_statistics(result,records)


def markdown(results):
    out=['# PERM 历史独立预览','',BOUNDARY,'','认证占比 = (Certified + Certified-Expired) / (Certified + Certified-Expired + Denied)，撤回不计分母；日期按纯日期，处理时长只计有效非负差值。P25/P75 使用线性插值。','']
    def cell(value):return str(value if value is not None else '不可得').replace('|','\\|').replace('\n',' ')
    for r in results:
        out += [f'## {r["company"]}','',f'匹配：{r["match_status"]}；版本：{r["revision"]}。','']
        out += ['覆盖：'+ '；'.join(f'FY{s["fiscal_year"]} Q{s["quarter"]}（{s["start"]} 至 {s["end"]}）' for s in r['coverage'])+'。','']
        if r.get('alias'):
            alias=r['alias'];out += [f'名称映射：{alias["source_company"]} → '+ ' / '.join(t['legal_name'] for t in alias['targets'])+f'；关系：{alias["relationship_type"]}；核实状态：{alias["status"]}；范围：{alias["scope"]["type"]}。','']
            if alias['scope']['type']=='posting':out += ['仅核实岗位：'+ '；'.join(p['posting_key']+' / '+p['posting_url'] for p in alias['scope']['postings'])+'。','']
            for e in alias['evidence']:out += [f'- 证据 [{e["id"]}]({e["url"]})（核实 {e["checked_at"]}；本地 {e["path"]}）：'+e['quote'],'']
        for e in r['matched_employers']:out += [f'官方法人：{e["name"]}（{e["city"]}, {e["state"]}）。','']
        if r['match_status']=='verified_alias' and len(r['entities'])>1:
            out += [r['match_reason'],'']
            for entity in r['entities']:
                separate={**r,**entity,'company':r['company']+' → '+entity['legal_name'],'entities':[entity],'matched_employers':[{'name':entity['legal_name'],**p} for p in entity['locations']]}
                separate.pop('alias',None);out += [markdown([separate])]
            continue
        if r['match_status'] not in ('exact','verified_alias'):
            out += [r.get('match_reason','无确认统计。名称候选不能证明同一法人。'),'',*[f'- {s["employer"]} / {s["trade"]}：仅名称候选' for s in r['suggestions']],''];continue
        out += ['|裁定财年|已披露|Certified|Certified-Expired|Denied|Withdrawn|曾获认证占比|','|---|---:|---:|---:|---:|---:|---|']
        for fy,s in r['years'].items():
            rate=s['historical_certification_share'];counts=s['status_counts'];pct=f'{rate["value"]:.1%}' if rate['value'] is not None else '不可计算'
            out.append('|'+ '|'.join(map(cell,[f'FY{fy}',s['disclosed_cases'],*[counts[k] for k in STATUSES],f'{rate["numerator"]}/{rate["denominator"]} ({pct})']))+'|')
        out += ['', '处理时长（天，合并所选披露期）：','','|结果|有效 n|排除|中位数|P25|P75|','|---|---:|---:|---:|---:|---:|']
        for status,s in r['processing_days'].items():out.append('|'+ '|'.join(map(cell,[status,s['n'],s['excluded'],s['median'],s['p25'],s['p75']]))+'|')
        out += ['',f'最近公开裁定：{cell(r["recent_dates"]["decision"])}；最近认证裁定：{cell(r["recent_dates"]["certification_decision"])}；披露记录中最近收到：{cell(r["recent_dates"]["received_among_disclosed"])}。','']
        for dimension,groups in r['breakdowns'].items():
            out += [f'### {dimension}','', '|分组（原文）|案件数|','|---|---:|']
            for g in groups:out.append('|'+cell(' / '.join(g['values'].values()))+'|'+str(g['disclosed_cases'])+'|')
            out.append('')
        out += ['### 案件证据（原表位置）','','|案件|状态|收到日|裁定日|源文件|Sheet|行|','|---|---|---|---|---|---|---:|']
        for c in r['cases']:out.append('|'+ '|'.join(map(cell,[c['case_number'],c['status'],c['received'],c['decision'],c['source_file'],c['sheet'],c['row']]))+'|')
        out.append('')
    return '\n'.join(out)+'\n'


def preview(path,ident):
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,79}',ident) or '..' in ident:raise ValueError('Invalid preview id')
    selection=read_json(path);companies=selection.get('companies',[])
    if not companies or len({norm(x['company']) for x in companies})!=len(companies):raise ValueError('Selection needs distinct scored companies')
    for c in companies:
        if not c.get('scored_evidence'):raise ValueError('Scored evidence required')
        for e in c['scored_evidence']:
            f=guarded(e['path'],inside=False)
            if not f.is_file():raise ValueError('Scored evidence file is missing')
    validation=selection.get('selection_validation')
    if validation:
        proof=guarded(validation['path'])
        if not proof.is_file():raise ValueError('Selection validation artifact is missing')
    db,manifest,pointer=load_current()
    try:
        aliases=load_aliases(db)
        results=[query(db,manifest,pointer,c['company'],aliases,[{'posting_key':e.get('posting_key'),'posting_url':e.get('posting_url')} for e in c['scored_evidence']]) for c in companies]
    finally:db.close()
    output=guarded(DATA/'previews'/ident);output.mkdir(parents=True,exist_ok=True)
    if any(output.iterdir()):raise ValueError('Preview already exists; choose a new id')
    alias_snapshot=None
    if aliases:
        frozen=json.loads(json.dumps(aliases['registry']));evidence_files=[]
        for ai,alias in enumerate(frozen['aliases']):
            for ei,e in enumerate(alias['evidence']):
                source=guarded(DATA/e['path']);relative=f'alias-evidence/{ai}-{ei}.txt';dest=guarded(output/relative);dest.parent.mkdir(exist_ok=True)
                shutil.copyfile(source,dest)
                e['source_path']=e['path'];e['path']=str(dest.relative_to(DATA));evidence_files.append({'path':relative})
        write_json(output/'employer-aliases.json',frozen)
        for r in results:
            if r.get('alias'):r['alias']=next(a for a in frozen['aliases'] if a['id']==r['alias']['id'])
        alias_snapshot={'path':'employer-aliases.json','evidence':evidence_files}
    write_json(output/'selection.json',selection)
    write_json(output/'preview.json',{'schema_version':1,'generated_at':now(),'dataset':pointer,'companies':results,'boundary':BOUNDARY})
    guarded(output/'preview.md').write_text(markdown(results))
    check_rows=[]
    for r in results:
        for status in STATUSES:
            for fy in (2025,2026):
                c=next((c for c in r.get('cases',[c for entity in r.get('entities',[]) for c in entity.get('cases',[])]) if c['status']==status and c['decision_fiscal_year']==fy),None)
                if c:check_rows.append({'company':r['company'],'status':status,'fiscal_year':fy,'case_number':c['case_number'],'source_file':c['source_file'],'sheet':c['sheet'],'row':c['row'],'manual_verification':'pending'})
    write_json(output/'spot-checks.json',{'checks':check_rows,'note':'Source locators for independent human/raw-workbook cross-check; not claimed as manually verified.'})
    write_json(output/'manifest.json',{'schema_version':1,'dataset':pointer,'company_count':len(results),'no_report_backfill':True,'evidence_validation':'File presence, source quotes and posting scope; canonical scoring/company identity validated by selection producer','selection_validation':validation,'alias_snapshot':alias_snapshot})
    for f in output.rglob('*'):
        if f.is_file():f.chmod(0o444)
    return {'status':'COMPLETE','directory':str(output),'company_count':len(results),'match_status_counts':{s:sum(r['match_status']==s for r in results) for s in sorted({r['match_status'] for r in results})}}


def scan_snapshot():
    db,manifest,pointer=load_current()
    try:
        aliases=load_aliases(db)
        registry=aliases['registry'] if aliases else {'schema_version':1,'aliases':[]}
        evidence={e['path']:guarded(DATA/e['path']).read_text() for a in registry['aliases'] for e in a['evidence']}
        # Recheck mutable metadata after copying; an in-progress refresh cannot mix versions.
        if read_json(DATA/'current.json')!=pointer or (read_json(DATA/'employer-aliases.json') if (DATA/'employer-aliases.json').exists() else {'schema_version':1,'aliases':[]})!=registry:
            raise ValueError('PERM source changed during baseline capture')
        return {'schema_version':1,'status':'ready','pointer':pointer,'manifest':manifest,'aliases':registry,'evidence':evidence}
    finally:db.close()


def scan_batch(payload):
    snapshot=payload['snapshot'];db,manifest,pointer=load_current(snapshot['pointer'])
    try:
        if snapshot.get('schema_version')!=1 or snapshot.get('status')!='ready' or manifest!=snapshot['manifest']:
            raise ValueError('Frozen PERM manifest changed; prepare a new run')
        registry=validate_aliases(snapshot['aliases'],db,snapshot['evidence'])
        references=[];postings=[];cache={};keys=set();suggestion_cache={}
        for record in payload['records']:
            key=record['posting_key'];url=record['posting_url'];company=record['company']
            if not key or key in keys or not url or not company:raise ValueError('PERM needs unique exact posting key, URL and company')
            keys.add(key)
            alias=next((a for a in registry['aliases'] if norm(a['source_company'])==norm(company)),None)
            group=(company,key,url) if alias and alias['scope']['type']=='posting' else (company,)
            if group not in cache:
                ident=f'perm-{len(references)+1}'
                result=query(db,manifest,pointer,company,{'registry':registry},[{'posting_key':key,'posting_url':url}],suggestion_cache)
                references.append({'id':ident,'result':result});cache[group]=ident
            postings.append({**record,'reference_id':cache[group]})
        lines=['# Daily Scan PERM 参考信息','',BOUNDARY,'','本页覆盖全部已评分岗位，也保留评分失败岗位的绑定；不改变评分、筛选或报告门槛。','','|公司 / 岗位键|评分成功|匹配|明细|','|---|---|---|---|']
        by_id={r['id']:r['result'] for r in references}
        for p in postings:
            name=p['company'].replace('|','\\|').replace('\n',' ')
            lines.append(f'|{name} / {p["posting_key"]}|{p["scored"]}|{by_id[p["reference_id"]]["match_status"]}|[详情](#{p["reference_id"]})|')
        for ref in references:lines+=['',f'<a id="{ref["id"]}"></a>','',markdown([ref['result']])]
        return {'schema_version':1,'dataset':pointer,'coverage':manifest['sources'],'postings':postings,'references':references,'markdown':'\n'.join(lines)+'\n'}
    finally:db.close()


def company_review_queue(payload):
    batch=scan_batch(payload);by_id={r['id']:r['result'] for r in batch['references']};companies={}
    for posting in batch['postings']:
        result=by_id[posting['reference_id']]
        if not posting['scored'] or result['match_status'] in ('exact','verified_alias') or result.get('alias_verification')=='verified':continue
        company=posting['company']
        companies.setdefault(company,{'company':company,'match_status':result['match_status'],'postings':[]})['postings'].append({'posting_key':posting['posting_key'],'posting_url':posting['posting_url']})
    return {'schema_version':1,'status':'required' if companies else 'not_required','companies':list(companies.values())}


def company_review(payload,freeze=False):
    baseline=payload['snapshot'];queue=company_review_queue(payload);review=payload.get('review',{});decisions=payload.get('decisions',review.get('companies',[]))
    if not isinstance(decisions,list):raise ValueError('Company review needs a decisions list')
    expected={c['company'] for c in queue['companies']};seen=set();verified=set()
    for decision in decisions:
        company=decision.get('company');outcome=decision.get('outcome')
        if company not in expected or company in seen:raise ValueError('Unexpected/duplicate company review decision')
        seen.add(company)
        if outcome not in ('verified','ambiguous','not_found','source_unavailable','publisher'):raise ValueError('Invalid company review outcome')
        if not isinstance(decision.get('note'),str) or not decision['note'].strip() or len(decision['note'])>2000:raise ValueError('Company review needs a bounded research note')
        if outcome=='ambiguous' and (not isinstance(decision.get('question'),str) or not decision['question'].strip() or len(decision['question'])>2000 or decision.get('question_state') not in ('asked','answered','answered_unknown')):raise ValueError('Ambiguous company requires a recorded user question')
        if outcome=='verified':verified.add(norm(company))
    if seen!=expected:raise ValueError('Company review is incomplete')
    snapshot=dict(baseline);aliases={norm(a['source_company']):a for a in baseline['aliases']['aliases']};evidence=dict(baseline['evidence'])
    if freeze:
        db,_,_=load_current(baseline['pointer'])
        try:updates=load_aliases(db)['registry']['aliases'] if verified else []
        finally:db.close()
    else:
        if review.get('schema_version')!=1 or review.get('status')!='COMPLETE' or review.get('records')!=payload['records']:raise ValueError('Frozen company review binding changed')
        updates=review['snapshot']['aliases']['aliases']
    by_name={norm(a['source_company']):a for a in updates}
    if not verified<=by_name.keys():raise ValueError('Company review needs a published verified company-history mapping')
    # Preserve registry/frozen order across Python processes with different hash seeds.
    for name in by_name:
        if name not in verified:continue
        alias=by_name.get(name)
        if not alias or alias['status']!='verified' or alias['relationship_type']!='company_history_reference' or alias['scope']['type']!='company':raise ValueError('Company review needs a published verified company-history mapping')
        aliases[name]=alias
        for item in alias['evidence']:
            value=guarded(DATA/item['path']).read_text() if freeze else review['snapshot']['evidence'].get(item['path'])
            if item['path'] in evidence and evidence[item['path']]!=value:raise ValueError('Company review cannot reinterpret baseline evidence; save a new evidence path')
            evidence[item['path']]=value
    snapshot.update(aliases={'schema_version':1,'aliases':list(aliases.values())},evidence=evidence)
    if not freeze and snapshot!=review['snapshot']:raise ValueError('Company review changed unrelated aliases, evidence or disclosure version')
    # The original disclosure revision and exact posting bindings remain authoritative.
    batch=scan_batch({'snapshot':snapshot,'records':payload['records']})
    result={'schema_version':1,'status':'COMPLETE','records':payload['records'],'companies':decisions,'snapshot':snapshot}
    return result if freeze else batch


def main():
    p=argparse.ArgumentParser(description=__doc__);commands=p.add_subparsers(dest='command',required=True)
    commands.add_parser('refresh')
    commands.add_parser('scan-snapshot');commands.add_parser('scan-batch')
    for command in ('company-review-queue','company-review-freeze','company-review-check'):commands.add_parser(command)
    i=commands.add_parser('import');i.add_argument('--manifest',required=True)
    a=commands.add_parser('aliases-import');a.add_argument('--file',required=True);a.add_argument('--merge',action='store_true')
    q=commands.add_parser('query');q.add_argument('--company',required=True);q.add_argument('--posting-key');q.add_argument('--posting-url')
    v=commands.add_parser('preview');v.add_argument('--companies',required=True);v.add_argument('--id',default=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%SZ'))
    args=p.parse_args();guarded(DATA)
    if args.command=='scan-snapshot':result=scan_snapshot()
    elif args.command=='scan-batch':result=scan_batch(without_checksums(json.load(sys.stdin)))
    elif args.command=='company-review-queue':result=company_review_queue(without_checksums(json.load(sys.stdin)))
    elif args.command in ('company-review-freeze','company-review-check'):result=company_review(without_checksums(json.load(sys.stdin)),freeze=args.command=='company-review-freeze')
    elif args.command=='refresh':
        with writer():
            page,sources=discover();download=guarded(DATA/'downloads'/('refresh-'+uuid.uuid4().hex));download.mkdir(parents=True)
            (download/'official-index.html').write_text(page)
            for s in sources:
                for key in ('xlsx','layout'):
                    url=s.pop(key+'_url');name=f'FY{s["fiscal_year"]}_Q{s["quarter"]}'+('.xlsx' if key=='xlsx' else '.pdf');s[key]=fetch(url,download/name)
            manifest={'schema_version':1,'official_index_url':INDEX,'discovered_at':now(),'sources':sources};write_json(download/'sources.json',manifest);result=import_sources(manifest)
    elif args.command=='import':
        with writer():result=import_sources(read_json(args.manifest))
    elif args.command=='aliases-import':
        with writer():result=import_aliases(args.file,args.merge)
    elif args.command=='query':
        db,manifest,pointer=load_current()
        try:result=query(db,manifest,pointer,args.company,load_aliases(db),[{'posting_key':args.posting_key,'posting_url':args.posting_url}] if args.posting_key or args.posting_url else [])
        finally:db.close()
    else:result=preview(args.companies,args.id)
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    try:main()
    except Exception as e:
        print(json.dumps({'status':'FAILED','error':str(e)},ensure_ascii=False),file=sys.stderr);sys.exit(1)
